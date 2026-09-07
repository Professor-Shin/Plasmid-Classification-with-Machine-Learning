from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "project_recap_for_advisor_v2.pptx"

NAVY = RGBColor(24, 42, 66)
BLUE = RGBColor(45, 111, 163)
TEAL = RGBColor(36, 132, 126)
ORANGE = RGBColor(211, 119, 54)
RED = RGBColor(171, 69, 69)
INK = RGBColor(35, 41, 48)
MUTED = RGBColor(92, 104, 115)
PALE = RGBColor(242, 246, 248)
WHITE = RGBColor(255, 255, 255)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", margin=0.06):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=18, color=INK, level_indent=0.25):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.space_after = Pt(8 if level == 0 else 3)
        p.font.name = "Aptos"
        p.font.size = Pt(size - (2 if level else 0))
        p.font.color.rgb = color
        p.font.bold = False
    return box


def base_slide(title, section=None, accent=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.16))
    band.fill.solid()
    band.fill.fore_color.rgb = accent
    band.line.fill.background()
    add_text(slide, title, 0.55, 0.42, 11.8, 0.55, size=27, color=NAVY, bold=True)
    if section:
        add_text(slide, section.upper(), 10.6, 0.48, 2.1, 0.3, size=9, color=accent, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, "Project recap | Plasmid classification", 0.58, 7.14, 4.5, 0.2, size=8, color=MUTED)
    add_text(slide, str(len(prs.slides)), 12.35, 7.12, 0.35, 0.2, size=8, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, title, body, x, y, w, h, color=BLUE, body_size=16):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = color
    shape.line.width = Pt(1.2)
    add_text(slide, title, x + 0.16, y + 0.12, w - 0.32, 0.32, size=16, color=color, bold=True)
    add_text(slide, body, x + 0.16, y + 0.52, w - 0.32, h - 0.62, size=body_size, color=INK)


def metric(slide, value, label, x, y, w, color=BLUE):
    add_text(slide, value, x, y, w, 0.5, size=25, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, y + 0.48, w, 0.42, size=11, color=MUTED, align=PP_ALIGN.CENTER)


# 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.0), Inches(0.12), Inches(5.1))
accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
add_text(slide, "Project Recap", 1.1, 1.3, 10.8, 0.6, size=34, color=WHITE, bold=True)
add_text(slide, "Plasmid Classification on DNA Sequence", 1.1, 2.05, 10.8, 0.65, size=29, color=WHITE, bold=True)
add_text(slide, "สถานะปัจจุบัน, สมมติฐาน, ข้อจำกัด และแผนการทดลองต่อไป", 1.1, 3.0, 10.6, 0.5, size=21, color=RGBColor(210, 224, 232))
add_text(slide, "สำหรับคุยกับอาจารย์ที่ปรึกษา | 7 กันยายน 2026", 1.1, 5.65, 8, 0.35, size=14, color=RGBColor(180, 198, 210))

# 2
slide = base_slide("1. Research question and motivation", "Context", TEAL)
card(slide, "Question", "Can a model distinguish plasmid from chromosome sequence in bacterial genomes?", 0.7, 1.35, 5.8, 1.55, TEAL, 20)
card(slide, "Why it matters", "Plasmids are important carriers of antimicrobial-resistance genes and can spread through horizontal gene transfer.", 6.85, 1.35, 5.8, 1.55, TEAL, 18)
add_bullets(slide, [
    "Target: a fast, alignment-free classifier for plasmid / chromosome",
    "Input: DNA sequence or assembled contig",
    "Output: binary class and evaluation metrics",
    "Research question: does contextual protein representation improve over k-mer baselines?",
], 1.0, 3.35, 11.5, 2.4, 20)

# 3
slide = base_slide("2. Current progress at a glance", "Status", BLUE)
card(slide, "Completed", "Dataset exploration\nK-mer + ML baselines\nExternal baseline runs\nArchitecture prototypes", 0.7, 1.35, 3.8, 3.8, TEAL, 18)
card(slide, "In progress", "Protein tokenization\nESM-2 embedding pipeline\nAttention pooling design\nFair comparison protocol", 4.8, 1.35, 3.8, 3.8, BLUE, 18)
card(slide, "Not yet established", "Attention + MLP on the full dataset\nGeneralization across species\nLeakage-safe fragment evaluation\nFinal proposed model", 8.9, 1.35, 3.8, 3.8, ORANGE, 18)
add_text(slide, "Key message: the project has a working baseline and a prototype architecture; the proposed method is not yet validated at scale.", 1.0, 5.75, 11.4, 0.65, size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

# 4
slide = base_slide("3. Data and evaluation setup", "Data", TEAL)
card(slide, "Main benchmark", "PlasmidHunter benchmark\n12,372 plasmids + 12,372 chromosomes\n1 kb to 100 kb\nMedian length: 10 kb", 0.7, 1.25, 3.75, 3.0, TEAL, 17)
card(slide, "Development data", "NCBI RefSeq-derived data\nCurrent prototype: 250 plasmids + 250 chromosomes\nUsed for training experiments", 4.8, 1.25, 3.75, 3.0, BLUE, 17)
card(slide, "Small prototype", "50 sequences per class\n20 held-out test sequences\nSingle random split\nCPU-constrained experiments", 8.9, 1.25, 3.75, 3.0, ORANGE, 17)
add_text(slide, "Important: the slide deck and repository contain multiple dataset protocols. Every final comparison must name the exact split and sample count.", 0.95, 5.0, 11.5, 0.7, size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)

# 5
slide = base_slide("4. Baseline pipeline and results", "Evidence", BLUE)
add_text(slide, "Canonical 5-mer frequency (512 features) -> classifier", 0.8, 1.15, 11.8, 0.4, size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
headers = ["Model", "Accuracy", "F1", "AUC"]
rows = [
    ("Logistic Regression", "0.6260", "0.6619", "0.6754"),
    ("Linear SVM", "0.7195", "0.7272", "0.7831"),
    ("Random Forest", "0.8717", "0.8706", "0.9501"),
    ("MLP", "0.8586", "0.8523", "0.9272"),
]
x0, y0 = 1.1, 1.9
widths = [4.2, 2.0, 2.0, 2.0]
for j, h in enumerate(headers):
    add_text(slide, h, x0 + sum(widths[:j]), y0, widths[j], 0.38, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + sum(widths[:j])), Inches(y0), Inches(widths[j]), Inches(0.42))
    rect.fill.solid(); rect.fill.fore_color.rgb = NAVY; rect.line.fill.background()
    rect._element.getparent().remove(rect._element); slide.shapes._spTree.insert(2, rect._element)
for i, row in enumerate(rows):
    yy = y0 + 0.45 + i * 0.65
    for j, value in enumerate(row):
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + sum(widths[:j])), Inches(yy), Inches(widths[j]), Inches(0.62))
        rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(235, 241, 245) if i % 2 == 0 else WHITE
        rect.line.color.rgb = RGBColor(210, 220, 226)
        add_text(slide, value, x0 + sum(widths[:j]), yy + 0.12, widths[j], 0.3, size=16, color=INK, bold=(i == 2), align=PP_ALIGN.CENTER)
metric(slide, "0.8706", "best F1: Random Forest", 2.0, 5.4, 3.0, TEAL)
metric(slide, "13 s", "RF training time", 5.2, 5.4, 3.0, BLUE)
metric(slide, "266 s", "MLP training time", 8.4, 5.4, 3.0, ORANGE)

# 6
slide = base_slide("5. External and literature baselines", "Evidence", ORANGE)
add_text(slide, "แยกจาก our trained baselines: ผลเหล่านี้ต้องใช้ test set และ evaluation rule เดียวกันก่อนจึงจะจัดอันดับได้", 0.75, 1.05, 11.9, 0.45, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
card(slide, "PlasClass", "Status: รันจริงใน repo\nF1 = 0.8588\nLightweight pretrained tool\nResult: results/pretrained_baselines/", 0.65, 1.75, 3.0, 3.7, TEAL, 16)
card(slide, "DeepLasmid", "Status: รันจริงผ่าน Docker\nF1 = 0.8490\nPrecision 0.9917 / Recall 0.7422\n963 predictions parsed", 3.9, 1.75, 3.0, 3.7, BLUE, 16)
card(slide, "mlplasmids", "Status: มี runner code แต่ไม่มี reproducible output\nSlide reports F1 = 0.5498\nNeeds Docker/R environment\nSpecies-specific concern", 7.15, 1.75, 3.0, 3.7, ORANGE, 16)
card(slide, "PlasmidHunter / PLASMe", "Status: ยังไม่ได้รันใน workspace\nAlignment/database heavy\nRequires external tools and large DB\nNot a weakness claim", 10.4, 1.75, 2.3, 3.7, RED, 15)
add_text(slide, "Takeaway: external methods are not absent because they are weak; the current gap is reproducibility, dependencies, and protocol mismatch.", 0.9, 5.95, 11.5, 0.55, size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)

# 7
slide = base_slide("6. What the baseline tells us", "Evidence", TEAL)
add_bullets(slide, [
    "Non-linear models are stronger than linear models on this benchmark.",
    "Random Forest is the current primary baseline: F1 = 0.8706.",
    "Performance depends strongly on fragment length:",
    ("1 kb: RF F1 = 0.6473", 1),
    ("10 kb: RF F1 = 0.8421", 1),
    ("100 kb: RF F1 = 0.9026", 1),
    "The baseline is strong enough that the proposed method must be compared against it under the same protocol.",
], 0.9, 1.25, 7.1, 4.9, 19)
card(slide, "Interpretation", "Short fragments have sparse/noisy k-mer profiles. This is the main opportunity for a contextual model, but it is also where leakage and split design matter most.", 8.35, 1.45, 4.1, 2.65, ORANGE, 18)
card(slide, "External reference", "PlasClass F1 = 0.8588\nDeepLasmid F1 = 0.8490\nNot directly comparable yet: protocols differ.", 8.35, 4.35, 4.1, 1.9, BLUE, 17)

# 8
slide = base_slide("7. Proposed method", "Architecture", BLUE)
card(slide, "Protein/function branch", "DNA\n-> ORF extraction\n-> amino-acid tokenization\n-> ESM-2 embedding\n-> attention pooling", 0.75, 1.4, 3.75, 3.6, TEAL, 18)
card(slide, "DNA/structure branch", "DNA\n-> 6-mer tokenization\n-> DNABERT windows\n-> attention pooling", 4.8, 1.4, 3.75, 3.6, BLUE, 18)
card(slide, "Fusion + classifier", "Concatenate contextual vectors\n-> MLP classifier\nCompare with:\n- classical ML\n- feed-forward MLP\n- attention + MLP", 8.85, 1.4, 3.75, 3.6, ORANGE, 17)
add_text(slide, "Current gap: embeddings and attention prototypes exist, but Attention + MLP has not yet been validated on the full, controlled dataset.", 1.0, 5.65, 11.3, 0.65, size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)

# 9
slide = base_slide("8. Assumptions to test, not conclusions", "Hypotheses", ORANGE)
card(slide, "H1: coding content", "If a large fraction of bacterial sequence is coding, translating ORFs to protein may expose functional signals more directly than raw k-mers.", 0.7, 1.25, 3.85, 3.8, ORANGE, 17)
card(slide, "H2: representation", "Pretrained ESM-2 embeddings may encode more useful biological context than k-mer counts, TF-IDF, or raw protein tokens.", 4.75, 1.25, 3.85, 3.8, BLUE, 17)
card(slide, "H3: aggregation", "Attention pooling may select informative ORFs instead of averaging all ORFs equally, but order/context must be defined and measured.", 8.8, 1.25, 3.85, 3.8, TEAL, 17)
add_text(slide, "Required evidence: ablation study + controlled split + confidence interval + comparison to RF 5-mer baseline.", 1.0, 5.75, 11.3, 0.55, size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

# 10
slide = base_slide("9. Current limitations", "Risks", RED)
add_bullets(slide, [
    "Only a small subset is used for Transformer training: typically 50 + / 50 -.",
    "Single random split and small test set make prototype scores unstable.",
    "Fragment evaluation can leak parent-sequence information across train/test.",
    "The comprehensive evaluator uses non-canonical k-mer extraction against a canonical vocabulary.",
    "ESM preprocessing drops short ORFs and non-coding regions; DNABERT uses limited windows.",
    "Some reported numbers use different protocols and are not directly comparable.",
], 0.95, 1.25, 11.5, 4.8, 19, RED)
add_text(slide, "Interpretation: the present results support feasibility, not final superiority.", 1.0, 6.1, 11.3, 0.45, size=21, color=RED, bold=True, align=PP_ALIGN.CENTER)

# 11
slide = base_slide("10. Proposed next experiment", "Plan", TEAL)
steps = [
    ("1", "Lock data manifest", "IDs, source, species, parent genome, label"),
    ("2", "Leakage-safe split", "Group by genome/species before fragments"),
    ("3", "Run common benchmark", "RF, SVM, MLP, PlasClass, DeepLasmid, ESM"),
    ("4", "Ablation", "Mean vs attention; ESM vs DNA vs hybrid"),
    ("5", "Scale on server", "Full training data, repeated runs, confidence intervals"),
]
for i, (num, title, body) in enumerate(steps):
    x = 0.75 + (i % 3) * 4.15
    y = 1.35 + (i // 3) * 2.35
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
    circ.fill.solid(); circ.fill.fore_color.rgb = TEAL; circ.line.fill.background()
    add_text(slide, num, x, y + 0.11, 0.62, 0.25, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.82, y + 0.02, 3.1, 0.3, size=17, color=NAVY, bold=True)
    add_text(slide, body, x + 0.82, y + 0.42, 3.15, 0.7, size=14, color=INK)
add_text(slide, "Decision point: validate the benchmark protocol before increasing model complexity.", 1.0, 6.25, 11.3, 0.4, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

# 12
slide = base_slide("11. Questions for advisor discussion", "Decision", ORANGE)
add_bullets(slide, [
    "Should the final contribution focus on ESM-2 protein representation, or keep the hybrid architecture as the main proposal?",
    "What is the accepted data split: random sequence split, species-level split, or genome-level split?",
    "Should the target be full sequences, simulated contigs, or both with separate evaluations?",
    "Which baseline set is mandatory: our RF/SVM/MLP, PlasClass, DeepLasmid, mlplasmids?",
    "Is the primary metric plasmid F1, macro-F1, recall, or a cost-sensitive metric for AMR screening?",
    "What result would count as a meaningful improvement over RF 5-mer F1 = 0.8706?",
], 0.85, 1.2, 11.7, 4.8, 18)
add_text(slide, "Goal of today: agree on the research question, evaluation protocol, and minimum experiment set.", 0.95, 6.25, 11.4, 0.45, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

# 13
slide = base_slide("Takeaway", "Close", TEAL)
add_text(slide, "What is already solid", 0.9, 1.25, 5.4, 0.4, size=22, color=TEAL, bold=True)
add_bullets(slide, ["Working dataset and code base", "Strong k-mer baseline", "Initial Transformer prototypes", "Clear engineering bottleneck"], 1.0, 1.85, 5.2, 2.7, 19)
add_text(slide, "What must happen next", 7.0, 1.25, 5.4, 0.4, size=22, color=ORANGE, bold=True)
add_bullets(slide, ["Make the split leakage-safe", "Run all models on one protocol", "Train beyond 50 + / 50 -", "Test whether context adds value"], 7.1, 1.85, 5.2, 2.7, 19)
add_text(slide, "The proposed method is a testable hypothesis, not yet a validated winner.", 1.0, 5.35, 11.3, 0.7, size=25, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides)} slides")