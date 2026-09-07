from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "project_recap_siwak_ci.pptx"

# The Siwak System CI Guideline palette.
BG = RGBColor(248, 249, 250)       # Alabaster White
INK = RGBColor(32, 33, 36)         # Charcoal
MUTED = RGBColor(95, 99, 104)
LINE = RGBColor(232, 234, 237)
GREEN = RGBColor(0, 200, 83)       # Active Green
BLUE = RGBColor(26, 115, 232)      # Ocean Blue
ORANGE = RGBColor(255, 87, 34)     # Warm Orange
YELLOW = RGBColor(251, 188, 5)     # Sunny Yellow
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def shape(slide, kind, x, y, w, h, fill=None, line=None, radius=False):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
         font="Noto Sans Thai", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_title(slide, black, accent, x=0.7, y=0.72, w=12.0, size=29):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.65))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    for value, color in [(black, INK), (accent, GREEN)]:
        run = p.add_run()
        run.text = value
        run.font.name = "Kanit"
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = color
    return box


def grid(slide):
    for x in range(0, 14):
        ln = shape(slide, MSO_SHAPE.RECTANGLE, x, 0, 0.006, 7.5, LINE)
        ln.fill.transparency = 20
    for y in range(0, 8):
        ln = shape(slide, MSO_SHAPE.RECTANGLE, 0, y, 13.333, 0.006, LINE)
        ln.fill.transparency = 20


def header(slide, section, number, accent=GREEN, title=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    grid(slide)
    # Small brand lockup, kept textual so the deck has no external asset dependency.
    text(slide, "●", 0.7, 0.29, 0.22, 0.25, size=12, color=GREEN, bold=True, font="Arial")
    text(slide, "SIWAK", 0.96, 0.25, 1.15, 0.28, size=12, color=INK, bold=True, font="Kanit")
    text(slide, section.upper(), 5.0, 0.3, 3.3, 0.24, size=9, color=MUTED, bold=True, align=PP_ALIGN.CENTER, font="Kanit")
    text(slide, f"{number:02d} / 10", 11.85, 0.28, 0.8, 0.24, size=10, color=BLUE, bold=True, align=PP_ALIGN.RIGHT, font="Kanit")
    shape(slide, MSO_SHAPE.RECTANGLE, 0.7, 0.62, 11.95, 0.012, LINE)
    if title:
        rich_title(slide, *title)
    text(slide, "PROJECT RECAP // ADVISOR DISCUSSION", 0.7, 7.12, 3.5, 0.16, size=7, color=MUTED, font="Kanit")


def card(slide, label, title, body, x, y, w, h, accent=BLUE, body_size=14):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.07, h, accent)
    text(slide, label.upper(), x + 0.25, y + 0.2, w - 0.45, 0.22, size=9, color=accent, bold=True, font="Kanit")
    text(slide, title, x + 0.25, y + 0.58, w - 0.45, 0.45, size=19, color=INK, bold=True, font="Kanit")
    text(slide, body, x + 0.25, y + 1.15, w - 0.45, h - 1.3, size=body_size, color=MUTED)


def bullets(slide, items, x, y, w, h, size=17, color=INK, accent=GREEN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(10)
        p.font.name = "Noto Sans Thai"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p._p.get_or_add_pPr().insert(0, p._p._new_buChar())
        p._p.pPr.insert(0, p._p._new_buChar())
    return box


def pill(slide, value, x, y, w, accent=GREEN):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.36, RGBColor(232, 248, 238), None)
    text(slide, value, x + 0.12, y + 0.07, w - 0.24, 0.2, size=10, color=accent, bold=True, font="Kanit", align=PP_ALIGN.CENTER)


def metric(slide, value, caption, x, y, w, accent=GREEN):
    text(slide, value, x, y, w, 0.5, size=27, color=accent, bold=True, font="Kanit", align=PP_ALIGN.CENTER)
    text(slide, caption, x, y + 0.48, w, 0.3, size=10, color=MUTED, align=PP_ALIGN.CENTER)


def table_cell(slide, value, x, y, w, h, fill, color=INK, bold=False, size=13):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, LINE)
    text(slide, value, x + 0.05, y + 0.09, w - 0.1, h - 0.12, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


# 1: cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
grid(slide)
text(slide, "●", 0.75, 0.45, 0.28, 0.28, size=15, color=GREEN, bold=True, font="Arial")
text(slide, "SIWAK", 1.08, 0.41, 1.3, 0.34, size=14, color=INK, bold=True, font="Kanit")
pill(slide, "ADVISOR DISCUSSION", 0.8, 1.3, 2.25, BLUE)
rich_title(slide, "PROJECT RECAP", "// NEXT EXPERIMENT", 0.8, 2.05, 11.5, 38)
text(slide, "Plasmid classification on DNA sequence", 0.85, 3.0, 7.8, 0.5, size=24, color=MUTED, font="Kanit")
text(slide, "สถานะที่ทำแล้ว  |  สมมติฐานที่ต้องทดสอบ  |  คำถามที่ต้องตัดสินใจ", 0.85, 3.65, 9.2, 0.35, size=17, color=INK)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.95, 1.6, 3.55, 3.55, WHITE, LINE)
shape(slide, MSO_SHAPE.ARC, 9.85, 2.35, 1.7, 1.7, None, BLUE)
shape(slide, MSO_SHAPE.ARC, 10.35, 2.85, 1.7, 1.7, None, GREEN)
shape(slide, MSO_SHAPE.OVAL, 11.18, 2.33, 0.22, 0.22, ORANGE)
text(slide, "DNA", 9.35, 4.27, 2.75, 0.35, size=17, color=BLUE, bold=True, font="Kanit", align=PP_ALIGN.CENTER)
text(slide, "EMBED  >  ATTEND  >  TEST", 0.85, 6.3, 5.0, 0.25, size=11, color=BLUE, bold=True, font="Kanit")
text(slide, "07 SEP 2026", 11.15, 6.3, 1.3, 0.25, size=10, color=MUTED, bold=True, font="Kanit", align=PP_ALIGN.RIGHT)

# 2: question
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "research question", 2, GREEN, ("WHAT ARE WE TRYING TO ", "ANSWER?"))
card(slide, "BIOLOGICAL PROBLEM", "Plasmid or chromosome?", "จำแนกลำดับ DNA ของแบคทีเรีย เพื่อช่วยศึกษาการแพร่ของ AMR ผ่าน HGT", 0.8, 1.65, 3.65, 2.45, GREEN, 16)
card(slide, "ENGINEERING GOAL", "Fast and alignment-free", "ลดการพึ่งพา database และเครื่องมือ alignment ที่หนักต่อ compute", 4.85, 1.65, 3.65, 2.45, BLUE, 16)
card(slide, "RESEARCH TEST", "Does context add value?", "วัดว่า protein embedding + attention ช่วยเหนือ k-mer baseline หรือไม่", 8.9, 1.65, 3.65, 2.45, ORANGE, 16)
text(slide, "คำถามหลักสำหรับงานต่อไป", 0.85, 4.85, 3.3, 0.28, size=14, color=GREEN, bold=True, font="Kanit")
text(slide, "เมื่อใช้ split และ test set เดียวกัน\ncontextual protein representation ช่วยเพิ่มประโยชน์จาก baseline ได้จริงหรือไม่?", 0.85, 5.25, 11.1, 0.75, size=23, color=INK, bold=True, font="Kanit")

# 3: progress
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "current status", 3, BLUE, ("WHAT IS ", "READY?"))
card(slide, "DONE", "Evidence base", "Dataset exploration\nK-mer + ML baselines\nPlasClass / DeepLasmid runs", 0.8, 1.55, 3.65, 3.25, GREEN, 17)
card(slide, "PROTOTYPE", "Architecture", "ESM-2 embeddings\nDNABERT branch\nAttention pooling\nHybrid fusion", 4.85, 1.55, 3.65, 3.25, BLUE, 17)
card(slide, "OPEN", "Validation gap", "Full-data Attention + MLP\nLeakage-safe split\nCommon benchmark\nFinal claim", 8.9, 1.55, 3.65, 3.25, ORANGE, 17)
pill(slide, "BASELINE READY", 0.9, 5.45, 1.9, GREEN)
pill(slide, "PROTOTYPE ONLY", 4.95, 5.45, 1.95, BLUE)
pill(slide, "DECISION NEEDED", 9.0, 5.45, 2.05, ORANGE)
text(slide, "สถานะที่ถูกต้อง: มี baseline ที่วัดได้แล้ว แต่ proposed method ยังไม่ใช่ผู้ชนะที่ validate แล้ว", 0.85, 6.25, 11.5, 0.35, size=16, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

# 4: data
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "data & protocol", 4, GREEN, ("NAME THE ", "PROTOCOL"))
text(slide, "ข้อมูลมีหลายระดับการทดลอง จึงต้องระบุ sample count ทุกครั้ง", 0.85, 1.42, 11.6, 0.3, size=17, color=MUTED)
card(slide, "FULL BENCHMARK", "12,372 + 12,372", "Plasmid + chromosome\n1 kb - 100 kb\nใช้กับ baseline หลัก", 0.8, 2.05, 3.65, 2.7, GREEN, 16)
card(slide, "DEVELOPMENT", "250 + 250", "NCBI RefSeq-derived\nใช้เตรียม training data\nสำหรับการทดลองขยาย", 4.85, 2.05, 3.65, 2.7, BLUE, 16)
card(slide, "PROTOTYPE", "50 + 50", "Attention log\nheld-out test 20\nsingle random split", 8.9, 2.05, 3.65, 2.7, ORANGE, 16)
text(slide, "ต้องล็อกก่อนรันรอบสุดท้าย", 0.85, 5.45, 2.8, 0.28, size=14, color=ORANGE, bold=True, font="Kanit")
text(slide, "source  /  parent genome  /  species group  /  train-test boundary  /  fragment policy", 0.85, 5.88, 11.4, 0.4, size=21, color=INK, bold=True, font="Kanit")

# 5: in-house results
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "our baseline", 5, BLUE, ("WHAT DOES THE ", "BASELINE SAY?"))
text(slide, "Canonical 5-mer frequency -> classifier", 0.85, 1.4, 5.8, 0.28, size=15, color=MUTED, font="Kanit")
x, y = 0.85, 1.95
widths = [3.35, 1.65, 1.65, 1.65]
for i, label in enumerate(["MODEL", "ACC", "F1", "AUC"]):
    table_cell(slide, label, x + sum(widths[:i]), y, widths[i], 0.45, INK, WHITE, True, 12)
rows = [("Logistic Regression", "0.6260", "0.6619", "0.6754"), ("Linear SVM", "0.7195", "0.7272", "0.7831"), ("Random Forest", "0.8717", "0.8706", "0.9501"), ("MLP", "0.8586", "0.8523", "0.9272")]
for r, row in enumerate(rows):
    for i, value in enumerate(row):
        table_cell(slide, value, x + sum(widths[:i]), y + 0.46 + r * 0.53, widths[i], 0.52, RGBColor(232, 248, 238) if r == 2 else WHITE, INK, r == 2, 12)
metric(slide, "0.8706", "best F1 / Random Forest", 8.25, 2.0, 3.6, GREEN)
metric(slide, "0.6473", "RF F1 / 1 kb", 8.25, 3.25, 3.6, ORANGE)
metric(slide, "0.9026", "RF F1 / 100 kb", 8.25, 4.5, 3.6, BLUE)
text(slide, "Short fragments are the opportunity area, but also the highest-risk area for leakage.", 0.85, 5.95, 11.2, 0.35, size=16, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

# 6: external
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "external baselines", 6, ORANGE, ("KEEP THE ", "COMPARISON HONEST"))
text(slide, "แยกจากโมเดลที่เราฝึกเอง: ตัวเลขยังมาจากคนละ protocol", 0.85, 1.38, 11.6, 0.28, size=16, color=MUTED)
x, y = 0.85, 1.9
widths = [2.55, 2.3, 1.45, 4.8]
for i, label in enumerate(["METHOD", "STATUS", "F1", "READ THIS AS"]):
    table_cell(slide, label, x + sum(widths[:i]), y, widths[i], 0.43, INK, WHITE, True, 11)
rows = [
    ("PlasClass", "รันจริง", "0.8588", "pretrained / lightweight"),
    ("DeepLasmid", "Docker run", "0.8490", "precision 0.9917, recall 0.7422"),
    ("mlplasmids", "ไม่มี raw output", "0.5498*", "ตัวเลขจาก slide; runner ต้องใช้ Docker/R"),
    ("PlasmidHunter / PLASMe", "ยังไม่รัน", "—", "alignment + large database; not a weakness claim"),
]
colors = [GREEN, BLUE, ORANGE, ORANGE]
for r, row in enumerate(rows):
    for i, value in enumerate(row):
        table_cell(slide, value, x + sum(widths[:i]), y + 0.44 + r * 0.72, widths[i], 0.71, WHITE, colors[r] if i == 1 else INK, i in (0, 1), 11)
text(slide, "* mlplasmids F1 = 0.5498 มีในสไลด์ แต่ยังไม่มี raw prediction / run log ใน repository", 0.9, 5.35, 11.5, 0.3, size=12, color=ORANGE)
text(slide, "Next: rerun all external tools on one locked test set before ranking them.", 0.9, 5.95, 11.5, 0.38, size=17, color=INK, bold=True, font="Kanit", align=PP_ALIGN.CENTER)

# 7: proposed method
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "proposed method", 7, BLUE, ("FROM DNA TO ", "CONTEXT"))
card(slide, "1 / INPUT", "DNA sequence", "raw sequence or contig", 0.75, 1.65, 2.25, 2.4, GREEN, 15)
card(slide, "2 / FUNCTION", "ORF -> protein", "amino-acid tokens\nESM-2 embedding", 3.25, 1.65, 2.55, 2.4, BLUE, 15)
card(slide, "3 / FOCUS", "Attention pooling", "learn which ORFs\nare informative", 6.05, 1.65, 2.55, 2.4, ORANGE, 15)
card(slide, "4 / DECISION", "MLP classifier", "plasmid / chromosome\ncompare to baselines", 8.85, 1.65, 3.35, 2.4, YELLOW, 15)
for x in [3.03, 5.83, 8.63]:
    text(slide, ">", x, 2.55, 0.18, 0.3, size=24, color=MUTED, bold=True, font="Kanit", align=PP_ALIGN.CENTER)
pill(slide, "OPTIONAL SECOND BRANCH: DNA -> DNABERT -> FUSION", 3.65, 4.75, 5.95, BLUE)
text(slide, "Current gap: Attention + MLP ยังไม่ได้ validate บน full data ด้วย protocol เดียวกับ RF", 0.85, 5.75, 11.5, 0.45, size=18, color=ORANGE, bold=True, font="Kanit", align=PP_ALIGN.CENTER)

# 8: hypotheses and limitations
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "hypotheses", 8, ORANGE, ("TEST THE ", "ASSUMPTIONS"))
card(slide, "H1", "Protein may expose function", "ถ้า coding signal สำคัญ การแปลง ORF เป็น protein อาจเห็น functional pattern ที่ k-mer ไม่เห็น", 0.8, 1.55, 3.65, 2.55, GREEN, 15)
card(slide, "H2", "ESM-2 may add context", "pretrained embedding อาจให้ biological context มากกว่า count หรือ raw token", 4.85, 1.55, 3.65, 2.55, BLUE, 15)
card(slide, "H3", "Attention may select signal", "attention อาจดีกว่า mean pooling ถ้า informative ORF มีเพียงบางส่วน", 8.9, 1.55, 3.65, 2.55, ORANGE, 15)
text(slide, "แต่ตอนนี้ยังมีความเสี่ยง", 0.85, 4.75, 2.8, 0.3, size=15, color=ORANGE, bold=True, font="Kanit")
text(slide, "50 + / 50 -   |   test = 20   |   single split   |   fragment leakage risk   |   protocol mismatch", 0.85, 5.2, 11.4, 0.4, size=20, color=INK, bold=True, font="Kanit")
text(slide, "ดังนั้นผล ESM-2 F1 = 0.7619 คือ preliminary evidence ไม่ใช่ final superiority", 0.85, 6.05, 11.5, 0.3, size=15, color=MUTED, align=PP_ALIGN.CENTER)

# 9: roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "next experiment", 9, GREEN, ("LOCK THE ", "EVIDENCE"))
steps = [("01", "Manifest", "ID / source / species / parent"), ("02", "Split", "group before fragment"), ("03", "Benchmark", "same test + same metrics"), ("04", "Ablation", "mean vs attention"), ("05", "Scale", "server + repeated runs")]
for i, (num, title, body) in enumerate(steps):
    x = 0.8 + i * 2.48
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 2.1, 2.5, WHITE, LINE)
    text(slide, num, x + 0.22, 2.25, 0.62, 0.3, size=21, color=GREEN if i < 3 else BLUE, bold=True, font="Kanit")
    text(slide, title, x + 0.22, 2.82, 1.65, 0.4, size=18, color=INK, bold=True, font="Kanit")
    text(slide, body, x + 0.22, 3.45, 1.65, 0.7, size=13, color=MUTED)
    if i < 4:
        text(slide, ">", x + 2.13, 3.0, 0.3, 0.3, size=20, color=MUTED, bold=True, font="Kanit", align=PP_ALIGN.CENTER)
text(slide, "Decision rule: validate the protocol before increasing model complexity.", 0.85, 5.55, 11.5, 0.4, size=21, color=GREEN, bold=True, font="Kanit", align=PP_ALIGN.CENTER)

# 10: decision
slide = prs.slides.add_slide(prs.slide_layouts[6]); header(slide, "advisor decision", 10, ORANGE, ("WHAT SHOULD WE ", "LOCK TODAY?"))
card(slide, "SCOPE", "ESM-only or hybrid?", "Protein branch first may be easier to validate; hybrid remains an ablation.", 0.8, 1.55, 3.65, 2.6, BLUE, 15)
card(slide, "SPLIT", "What is the unit?", "sequence, genome, species, or parent group? This controls leakage risk.", 4.85, 1.55, 3.65, 2.6, GREEN, 15)
card(slide, "SUCCESS", "What counts as better?", "F1, recall, PR-AUC, runtime, or a trade-off against RF F1 = 0.8706?", 8.9, 1.55, 3.65, 2.6, ORANGE, 15)
text(slide, "END STATE", 0.85, 5.05, 1.2, 0.25, size=11, color=BLUE, bold=True, font="Kanit")
text(slide, "ได้ research question + split + baseline set + minimum experiment", 0.85, 5.45, 11.2, 0.48, size=25, color=INK, bold=True, font="Kanit", align=PP_ALIGN.CENTER)
pill(slide, "BASELINE READY  /  PROPOSED METHOD TO VALIDATE", 3.75, 6.25, 5.85, GREEN)

prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides)} slides")